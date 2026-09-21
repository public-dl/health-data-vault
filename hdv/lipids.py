"""Audited source-column metadata; independent ratios, never a composition."""
from pathlib import Path
from openpyxl import load_workbook
from .common import read_json, checked_path, digest

MEMBERS = [('lipid_people', '実人員', 'R', '#124b9b'),
           ('triglycerides', '中性脂肪', 'S', '#54aeb4'),
           ('hdl', 'HDLコレステロール', 'T', '#629fce'),
           ('ldl', 'LDLコレステロール', 'U', '#54b5b1'),
           ('total_cholesterol', '総コレステロール', 'V', '#8e9bad')]
NOTE = '本データは原資料の「判定区分（保健指導以上を再掲）」の脂質代謝欄に掲載された人数です。割合は同年度・同地域の特定健診受診者数を分母として算出しています。各項目は排他的な構成ではないため、割合を合計して100%として扱うことはできません。実人員は原表名称のまま表示しています。具体的な重複排除・集計条件までは原資料から確定していません。'


RATE_MAP_BREAKS = {
    'lipid_people': [60,65,70,75],
    'triglycerides': [24,28,32,36],
    'hdl': [4,5,6,8],
    'ldl': [40,45,50,55],
    'total_cholesterol': [25,30,35,40],
}


def indicator_metadata(fixed_breaks=True):
    return [dict(indicator_id=id, name='脂質：'+label, group='lipids', theme_id='lipids',
        theme_label='脂質', unit='人', description='原資料の脂質代謝欄に掲載された人数。各項目は完全構成ではありません。',
        visualization_type='single_judgment_rate', overview_label='掲載項目をひと目で見る',
        count_label='報告人数', recipient_label='特定健診受診者数', display_set_label='脂質：全項目',
        set_notice='全項目は原資料に掲載された5項目の一覧です。各項目は排他的な構成ではないため、割合を合計して100%として扱うことはできません。',
        source_notice=NOTE, color=color, map_breaks=[100,500,2000,10000],
        intervals={'2021_2022':'pending','2022_2023':'pending'},
        rate=dict(label='特定健診受診者に占める割合（%）',map_breaks=list(RATE_MAP_BREAKS[id]) if fixed_breaks else [20,40,60,80],
                  intervals={'2021_2022':'pending','2022_2023':'pending'}))
        for id,label,_,color in MEMBERS]


def audit_formula_evidence(root, analysis, members=MEMBERS, official=None):
    """Read original workbooks, verifying formula text, cached value and source cells."""
    manifest = read_json(checked_path(root, 'processed/'+analysis['input_run_id']+'/collection.json'))
    result = {id:{} for id,_,_,_ in members}
    if official is None:
        official = {'triglycerides':(15,16), 'hdl':(17,18), 'ldl':(19,20)}
    for source in manifest['sources']:
        path = checked_path(root, source['raw_path'])
        if digest(path.read_bytes()) != source['sha256']:
            raise ValueError('Lipid source hash mismatch')
        formula = load_workbook(path, data_only=False, keep_links=False)
        cached = load_workbook(path, data_only=True, keep_links=False)
        try:
            overview, main = formula.worksheets[0], cached.worksheets[1]
            year = source['observation_fiscal_year']
            if year not in (2021,2022,2023): raise ValueError('Unaudited lipid year')
            if cached.worksheets[0]['J4'].value != main['B6'].value:
                raise ValueError('Lipid official denominator mismatch')
            for id,label,col,_ in members:
                evidence = dict(origin='hdv-derived', source_sha256=source['sha256'],
                    source_sheet=main.title, numerator_column=col, denominator_column='B',
                    rationale='原表人数÷同年度・同地域の特定健診受診者数×100。HDVで明示的に定義した独立割合。')
                if id in official:
                    nr,rr = official[id]; expected=f'=J{nr}/J'+('4' if year==2023 else '$4')+'*100'
                    if overview[f'J{rr}'].value != expected:
                        raise ValueError('Lipid official formula changed')
                    n=cached.worksheets[0][f'J{nr}'].value; d=main['B6'].value
                    if n != main[f'{col}6'].value or not d or abs(cached.worksheets[0][f'J{rr}'].value-n/d*100)>1e-9:
                        raise ValueError('Lipid official fraction mismatch')
                    evidence.update(origin='official-formula-confirmed', overview_sheet=overview.title,
                        numerator_cell=f'J{nr}', denominator_cell='J4', rate_cell=f'J{rr}', formula=expected,
                        rationale='県概要の同算式を確認。市町村割合は各市町村の原表人数とB列からHDVで算出。')
                result[id][str(year)] = evidence
        finally:
            formula.close(); cached.close()
    if any(set(v)!={'2021','2022','2023'} for v in result.values()):
        raise ValueError('Incomplete lipid formula evidence')
    return result
