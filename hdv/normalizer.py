"""Read-only conversion, explicit source cell states, conservative row selection."""
import math
import re
import uuid
from datetime import date, datetime
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .common import (CONFIG, checked_path, digest, issue, now, read_json,
                     write_json, write_jsonl)

MUNICIPAL_SCOPE = "新潟県内市町村国保40～74歳・国保連費用決済データから把握した健診受診者（法定報告とは異なる）"
UNION_SCOPE = "県内国保組合の報告受診者（市町村国保とは別集計・v0.1比較対象外）"


def structure(workbook):
    result = []
    for index, sheet in enumerate(workbook):
        fixed = {}
        for row in sheet:
            for cell in row:
                # Main table: freeze headings/row labels, not observed data or formulas.
                if cell.value is not None and (index != 1 or cell.row <= 5 or cell.column == 1):
                    if index == 1 or isinstance(cell.value, str):
                        fixed[cell.coordinate] = cell.value
        result.append(dict(name=sheet.title, rows=sheet.max_row, columns=sheet.max_column,
                           merged=sorted(str(r) for r in sheet.merged_cells.ranges),
                           print_area=str(sheet.print_area), fixed_cells=fixed))
    return result


def cell_state(cell, cached):
    raw, value = cell.value, cached.value
    formula = raw if cell.data_type == "f" else None
    if cell.data_type == "e" or cached.data_type == "e":
        state, numeric = "excel_error", None
    elif formula is not None and value is None:
        state, numeric = "formula_cache_missing", None
    elif value is None:
        state, numeric = "blank", None
    elif isinstance(value, bool):
        state, numeric = "unexpected_boolean", None
    elif isinstance(value, (int, float)):
        state, numeric = ("zero" if value == 0 else "numeric"), value
        if not math.isfinite(value):
            state, numeric = "invalid_number", None
    else:
        state = {"-": "hyphen", "－": "hyphen", "―": "hyphen", "・": "not_applicable_marker",
                 "": "empty_string", "未実施": "not_performed", "実施なし": "not_performed",
                 "未把握": "not_ascertained", "判定不能": "indeterminate_text"}.get(value, "unexpected_text")
        numeric = None
    def serializable(v):
        return v.isoformat() if isinstance(v, (date, datetime)) else v
    return dict(value=numeric, value_state=state, original_value=serializable(raw), cached_value=serializable(value),
                source_data_type=cell.data_type, number_format=cell.number_format, formula=formula)


def geography(name):
    compact = re.sub(r"\s", "", name)
    if compact == "県計":
        return compact, "prefecture_total", True
    if compact == "市計":
        return compact, "city_total", True
    if compact == "町村計":
        return compact, "town_village_total", True
    if compact.endswith("保健所計"):
        return compact, "health_center_total", True
    if compact == "国保組合計":
        return compact, "insurance_union_total", False
    if compact.endswith("国保"):
        return compact, "insurance_union", False
    if compact.endswith(("市", "町", "村")):
        return compact, "municipality", True
    return compact, "unknown", False


def select_rows(records):
    """Only the previously documented identical Niigata pair may be resolved automatically."""
    by_geo = defaultdict(lambda: defaultdict(list))
    for r in records:
        key = (r["observation_fiscal_year"], r["geography_level"], r["geography_name"])
        by_geo[key][r["source_row"]].append(r)
    decisions, issues = [], []
    for key, rows in by_geo.items():
        if len(rows) <= 1:
            continue
        row_numbers = sorted(rows)
        def signature(rs):
            return [(r["indicator_id"], r["original_value"], r["cached_value"],
                     r["source_data_type"], r["number_format"]) for r in rs]
        identical = all(signature(rows[n]) == signature(rows[row_numbers[0]]) for n in row_numbers[1:])
        known = key[2] == "新潟市" and row_numbers == [50, 51] and identical
        reason = ("Documented identical rows 50/51: retain row 50; preserve row 51 as excluded evidence."
                  if known else "Unexpected or conflicting duplicate: quarantine every occurrence; manual review required.")
        for n in row_numbers:
            for record in rows[n]:
                record["selection_reason"] = reason
                record["duplicate_group"] = f"{key[0]}:{key[1]}:{key[2]}"
                if not known or n != 50:
                    record["included"] = False
                    record["selection_status"] = "duplicate_excluded" if known else "duplicate_conflict"
        decisions.append(dict(observation_fiscal_year=key[0], geography_name=key[2],
                              source_rows=row_numbers, identical=identical,
                              retained_row=50 if known else None, reason=reason))
        issues.append(issue("documented_duplicate" if known else "unexpected_duplicate",
                            "warning" if known else "error", year=key[0], geography=key[2], rows=row_numbers))
    return decisions, issues


def extract(path, source, contract, indicators):
    workbook = load_workbook(path, data_only=False, keep_links=False)
    cached = load_workbook(path, data_only=True, keep_links=False)
    issues, cells, records = [], [], []
    year = source["observation_fiscal_year"]
    actual_structure = structure(workbook)
    structure_ok = actual_structure == contract["structure"]
    if not structure_ok:
        issues.append(issue("structure_changed", year=year, expected=contract["structure"], actual=actual_structure))
    # Preserve both sheets, including overview/formulas and every blank cell within used rectangles.
    for sheet in workbook:
        cs = cached[sheet.title]
        for row in sheet:
            for cell in row:
                state = cell_state(cell, cs[cell.coordinate])
                cells.append(dict(observation_fiscal_year=year, source_url=source["source_url"],
                                  source_sha256=source["sha256"], source_sheet=sheet.title,
                                  source_cell=cell.coordinate, **state))
                if state["value_state"] in {"excel_error", "formula_cache_missing"}:
                    issues.append(issue(state["value_state"], year=year, sheet=sheet.title, cell=cell.coordinate))
    if contract["main_sheet"] not in workbook.sheetnames:
        workbook.close()
        cached.close()
        return records, cells, issues, []
    sheet = workbook[contract["main_sheet"]]
    cs = cached[sheet.title]
    for row in range(6, sheet.max_row + 1):
        label = sheet.cell(row, 1).value
        if not isinstance(label, str):
            issues.append(issue("missing_geography", year=year, row=row))
            continue
        name, level, in_scope = geography(label)
        for indicator in indicators:
            column = indicator["column"]
            cell = sheet[f"{column}{row}"]
            state = cell_state(cell, cs[cell.coordinate])
            included = in_scope and structure_ok
            status = indicator["comparability_status"]
            if not in_scope:
                status = "pending"
            record = dict(
                schema_version="0.1", publication_fiscal_year=source["publication_fiscal_year"],
                observation_fiscal_year=year, geography_name=name, geography_original_label=label,
                geography_level=level, indicator_id=indicator["indicator_id"],
                original_label=" / ".join(str(sheet[c].value) for c in indicator["label_cells"]),
                source_label_cells=indicator["label_cells"], unit=indicator["unit"],
                population_scope=MUNICIPAL_SCOPE if in_scope else UNION_SCOPE,
                source_url=source["source_url"], annual_page_url=source["annual_page_url"],
                source_sheet=sheet.title, source_cell=cell.coordinate, source_row=row,
                source_sha256=source["sha256"], retrieved_at=source["retrieved_at"],
                comparability_status=status, comparability_intervals=indicator["intervals"],
                comparability_reason=indicator["reason"] if in_scope else "Outside municipal NHI scope",
                definition_version=f"niigata-tokutei-{year}-v0.1",
                definition_reference="docs/tokuteikenshin_definitions.md#52-対象excelをそのまま時系列接続する判断",
                comparison_group=indicator["indicator_id"] + ":reported-counts-r3-r5-v0.1",
                judgement_category="indeterminate" if column in ("F", "O") else None,
                validation_status="unvalidated", comparison_allowed=False, included=included,
                selection_status="included" if included else ("out_of_scope" if not in_scope else "structure_quarantined"),
                selection_reason="Primary municipal table" if included else (
                    "Insurance unions retained separately; not part of municipal totals" if not in_scope
                    else "Source structure changed; all source cells retained for review"),
                **state)
            records.append(record)
    decisions, duplicates = select_rows(records)
    issues.extend(duplicates)
    workbook.close()
    cached.close()
    return records, cells, issues, decisions


def normalize(root, manifest=None):
    root = Path(root)
    manifest = manifest or read_json(root / "raw" / "latest.json")
    contracts = read_json(CONFIG / "contracts.json")
    indicators = read_json(CONFIG / "indicators.json")
    records, cells, issues, decisions = [], [], [], []
    years = [s["observation_fiscal_year"] for s in manifest["sources"]]
    if sorted(years) != [2021, 2022, 2023]:
        raise ValueError("Collection must contain exactly observation years 2021, 2022, 2023")
    for source in manifest["sources"]:
        path = checked_path(root, source["raw_path"])
        if digest(path.read_bytes()) != source["sha256"]:
            raise ValueError("Raw SHA-256 mismatch")
        result = extract(path, source, contracts[str(source["observation_fiscal_year"])], indicators)
        records.extend(result[0]); cells.extend(result[1]); issues.extend(result[2]); decisions.extend(result[3])
    run_id = uuid.uuid4().hex
    output = root / "processed" / run_id
    write_jsonl(output / "records.jsonl", records)
    write_jsonl(output / "source_cells.jsonl", cells)
    write_json(output / "normalization_issues.json", issues)
    write_json(output / "row_decisions.json", decisions)
    write_json(output / "collection.json", manifest)
    write_json(output / "run.json", dict(run_id=run_id, created_at=now(), software_version="0.1.0",
                                        configuration_sha256={p.name: digest(p.read_bytes()) for p in CONFIG.glob("*.json")}))
    write_json(root / "processed" / "latest.json", {"run_id": run_id})
    return output
