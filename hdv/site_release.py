"""Independent annual/table contract; never relax the temporal public validator."""
import argparse
import copy
import json
from pathlib import Path

from openpyxl import load_workbook
from .common import CONFIG, atomic_write, checked_path, digest, now, read_json, write_json
from .normalizer import MUNICIPAL_SCOPE, cell_state, geography, structure
from .publisher import encoded, inspect as inspect_analysis, verified_records, locked
from .rates import usable, FORMULA

VERSION = 'physician-annual-v1'
REFERENCE = 'docs/annual_publication_design.md'
MEMBERS = [('normal', 'physician_normal', '異常認めず', '#54aeb4'),
           ('guidance', 'physician_guidance', '保健指導', '#d9a052'),
           ('referral', 'physician_referral', '受診勧奨', '#df806b')]


def require(condition, message):
    if not condition:
        raise ValueError(message)


def annual_payload(analysis, records):
    years = [2021, 2022, 2023]  # Explicit audited scope, never infer new years.
    require(analysis['years'] == years and len(analysis['geographies']) == 31, 'Unaudited annual scope')
    geos = {g['original_name']: g for g in analysis['geographies']}
    ids = [m[1] for m in MEMBERS]
    groups = {}
    for raw in records:
        if not raw['included'] or raw['geography_level'] not in ('municipality', 'prefecture_total'):
            continue
        if raw['indicator_id'] not in [*ids, 'recipients', 'physician_total']:
            continue
        r = copy.deepcopy(raw)
        require(r['geography_name'] in geos, 'Unknown annual region')
        r['geography_code'] = geos[r['geography_name']]['code']
        r['record_id'] = digest(encoded([r['source_sha256'], r['source_sheet'], r['source_cell']]))[:24]
        key = (r['observation_fiscal_year'], r['geography_code'])
        subset = groups.setdefault(key, {})
        require(r['indicator_id'] not in subset, 'Duplicate annual record')
        subset[r['indicator_id']] = r
    require(set(groups) == {(y, g['code']) for y in years for g in geos.values()}, 'Incomplete annual coverage')
    rows, denominators, proofs, insights = [], [], [], []
    for (year, code), subset in sorted(groups.items()):
        require(set(subset) == {*ids, 'recipients', 'physician_total'}, 'Missing annual component')
        d = subset['recipients']
        for r in subset.values():
            require(usable(r), 'Invalid annual numeric input')
            require(r['population_scope'] == MUNICIPAL_SCOPE, 'Annual population mismatch')
            for field in ('observation_fiscal_year', 'publication_fiscal_year', 'geography_name',
                          'geography_level', 'population_scope', 'source_sha256', 'source_sheet', 'source_row'):
                require(r[field] == d[field], 'Annual input mismatch: '+field)
        total = sum(subset[k]['value'] for k in ids)
        require(d['value'] > 0 and total == subset['physician_total']['value'] == d['value'], 'Annual partition mismatch/zero denominator')
        denominators.append(d)
        proofs.append(dict(group_id='doctor_judgment', observation_fiscal_year=year, geography_code=code,
                           category_sum=total, denominator_value=d['value'], difference=total-d['value'],
                           category_record_ids=[subset[k]['record_id'] for k in ids], denominator_record_id=d['record_id'],
                           total_record=subset['physician_total'], validation_status='passed', definition_version=VERSION))
        for id in ids:
            r = subset[id]
            require(r['comparability_status'] == 'pending' and not r['comparison_allowed'], 'Temporal policy changed; review required')
            r['annual_display_allowed'] = True
            r['derived_rate'] = dict(value=r['value']/d['value']*100, value_state='zero' if r['value']==0 else 'numeric',
                unit='%', numerator_value=r['value'], denominator_value=d['value'], numerator_record_id=r['record_id'],
                denominator_record_id=d['record_id'], formula=FORMULA, definition_version=VERSION,
                definition_reference=REFERENCE, comparability_status='pending', comparison_allowed=False,
                comparability_intervals=dict(r['comparability_intervals']), validation_status='passed', annual_display_allowed=True)
            rows.append(r)
            for measure in ('count', 'rate'):
                value = f"{r['value']:,}人" if measure=='count' else f"{r['derived_rate']['value']:.1f}%"
                insights.append(dict(indicator_id=id, geography_code=code, observation_fiscal_year=year, measure=measure,
                    text=f"{year}年度の{geos[r['geography_name']]['name']}では、この区分は{value}です。原表の報告区分であり、健康状態の優劣や疾病有病率を示しません。年度間の比較可能性は確認中です。", generator='annual-template-v1'))
    intervals = {'2021_2022':'pending', '2022_2023':'pending'}
    indicators, categories = [], []
    for order, (category, id, label, color) in enumerate(MEMBERS):
        maximum = max(r['value'] for r in rows if r['indicator_id']==id and r['geography_level']=='municipality')
        step = max(1, int((maximum+4)//5))
        breaks = [step*i for i in range(1,5)]
        indicators.append(dict(indicator_id=id, name='医師の判断：'+label, group='doctor_judgment', unit='人',
            comparability_status='pending', annual_display_allowed=True, comparison_allowed=False,
            description='原表で報告された医師の判断区分。単年度表示のみ。', map_breaks=breaks, intervals=intervals,
            rate=dict(label='受診者に占める割合（%）', map_breaks=[20,40,60,80], intervals=intervals)))
        categories.append(dict(category_id=category, indicator_id=id, label=label, color=color, order=order,
                               count_map_breaks=breaks, rate_map_breaks=[20,40,60,80]))
    group = dict(group_id='doctor_judgment', name='医師の判断', definition_version=VERSION, audit_reference=REFERENCE,
        composition_rule='reported_categories_sum_equals_recipients', denominator_indicator_id='recipients',
        required_category_ids=[m[0] for m in MEMBERS], allowed_views=['annual_composition','annual_map','annual_table'], categories=categories)
    return dict(schema_version='annual-1', input_run_id=analysis['input_run_id'], years=years, population_scope=MUNICIPAL_SCOPE,
        indicators=indicators, indicator_groups=[group], geographies=analysis['geographies'], records=rows,
        denominator_records=denominators, composition_validation=proofs, map=analysis['map'], insights=insights,
        source_validation=analysis['source_validation'], warnings=['年度間の比較可能性を確認中のため、現在は各年度を個別に表示しています。'])


def display_cell(state):
    value = state['cached_value']
    if value is None:
        return ''
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        require(state['number_format'] == '#,##0;\\-#,##0;\\-', 'Unknown source numeric display format')
        require(float(value).is_integer(), 'Unexpected fractional source value')
        return '-' if value == 0 else f'{value:,}'
    return str(value)


def published_tables(root, analysis, annual, reported=None):
    manifest = read_json(checked_path(root, 'processed/'+analysis['input_run_id']+'/collection.json'))
    contracts = read_json(CONFIG/'contracts.json')
    lookup = {(r['source_sha256'],r['source_sheet'],r['source_cell']):
              dict(indicator_id=r['indicator_id'], region=r['geography_code']) for r in analysis['records']+annual['records']+(reported['records'] if reported else [])}
    result = []
    for source in manifest['sources']:
        contract = contracts[str(source['observation_fiscal_year'])]
        path = checked_path(root, source['raw_path'])
        require(digest(path.read_bytes()) == source['sha256'], 'Source hash mismatch')
        book = load_workbook(path, data_only=False, keep_links=False)
        cached = load_workbook(path, data_only=True, keep_links=False)
        try:
            require(structure(book) == contract['structure'], 'Source structure changed')
            sheet, cs = book[contract['main_sheet']], cached[contract['main_sheet']]
            rows = []
            for row in sheet:
                cells = []
                for cell in row:
                    state = cell_state(cell, cs[cell.coordinate])
                    require(state['value_state'] not in ('excel_error','formula_cache_missing','invalid_number'), 'Unrenderable source cell')
                    cells.append(dict(row=cell.row, column=cell.column, coordinate=cell.coordinate, **state,
                        display_text=display_cell(state), visualization=lookup.get((source['sha256'],sheet.title,cell.coordinate))))
                rows.append(dict(index=row[0].row, level=geography(str(row[0].value))[1] if row[0].row>=6 else 'header', cells=cells))
            # Verify official overview explicitly labels this partition and publishes recipient-based fractions.
            overview = book.worksheets[0]
            require(overview['B5'].value == '医師の判断' and overview['C8'].value == '構成割合', 'Missing official composition evidence')
            for n, row in enumerate((8,9,10)):
                require(cached.worksheets[0].cell(row,10).value is not None, 'Missing official fraction cache')
                expected = cs.cell(6,8+n).value/cs['B6'].value*100
                require(abs(cached.worksheets[0].cell(row,10).value-expected)<1e-9, 'Official composition mismatch')
            result.append(dict(schema_version='published-table-1', source=source, source_sheet=sheet.title,
                header_rows=5, row_count=sheet.max_row, column_count=sheet.max_column,
                merged_ranges=[str(r) for r in sheet.merged_cells.ranges], rows=rows,
                composition_reference=dict(sheet=overview.title, cells=['B5','C8','J4','J5','J6','J7','J8','J9','J10'])))
        finally:
            book.close(); cached.close()
    return result


def construct(root, base_id, include_reported=False, reported_rates=False, include_lipids=False, lipid_fixed_breaks=False, include_glucose=False, unified_rates=False, terminology_version=None, include_liver=False, include_renal=False):
    analysis, report = inspect_analysis(root, base_id)
    records, _ = verified_records(root, analysis['input_run_id'])
    annual = annual_payload(analysis, records)
    from .reported_annual import reported_payload
    from .lipids import audit_formula_evidence
    evidence = audit_formula_evidence(root, analysis) if include_lipids else None
    from .glucose import audit_formula_evidence as audit_glucose
    glucose_evidence = audit_glucose(root, analysis) if include_glucose else None
    from .liver import audit_formula_evidence as audit_liver
    require(not include_liver or (include_glucose and unified_rates and terminology_version), 'liver requires common rate contract and terminology')
    liver_evidence = audit_liver(root, analysis) if include_liver else None
    reported = reported_payload(analysis, records, with_rates=reported_rates, lipid_evidence=evidence, lipid_fixed_breaks=lipid_fixed_breaks, glucose_evidence=glucose_evidence, liver_evidence=liver_evidence) if include_reported else None
    if unified_rates:
        from .reported_rate_contract import apply_contract
        reported = apply_contract(reported, include_liver=include_liver)
    if terminology_version is not None:
        from .terminology import apply_terminology
        reported = apply_terminology(reported, terminology_version)
    if include_renal:
        require(include_liver and unified_rates, 'renal requires verified current foundation')
        from .renal import extend
        reported = extend(reported, root, analysis, records)
    tables = published_tables(root, analysis, annual, reported)
    payload = dict(schema_version='site-1', analysis_release_id=base_id, analysis=analysis, annual=annual, published_tables=tables)
    if reported is not None:
        payload['reported'] = reported
    return payload, dict(reported_records=len(reported['records']) if reported else 0, errors=0, status='passed', analysis=report, annual_groups=len(annual['composition_validation']),
        annual_records=len(annual['records']), annual_comparability='pending', tables=len(tables),
        table_cells=sum(t['row_count']*t['column_count'] for t in tables))


def build(root, base_id, include_renal=False):
    from .terminology import configuration
    payload, report = construct(root, base_id, include_reported=True, reported_rates=True, include_lipids=True, lipid_fixed_breaks=True, include_glucose=True, unified_rates=True, terminology_version=configuration()['current_version'], include_liver=True, include_renal=include_renal)
    content = encoded(payload); id = digest(content)
    directory = Path(root)/'site/candidates'/id
    with locked(Path(root)/'site'):
        if not directory.exists():
            atomic_write(directory/'data.json', content)
            write_json(directory/'validation.json', report)
    return id, report


def inspect(root, id):
    require(isinstance(id,str) and len(id)==64 and all(c in '0123456789abcdef' for c in id), 'Invalid site release id')
    raw = (checked_path(root, 'site/candidates/'+id)/'data.json').read_bytes()
    require(digest(raw)==id, 'Site hash mismatch')
    payload = json.loads(raw)
    version = payload.get('reported',{}).get('schema_version')
    expected, report = construct(root, payload['analysis_release_id'], include_reported='reported' in payload, reported_rates=version in ('reported-annual-2','reported-annual-3','reported-annual-4','reported-annual-5'), include_lipids=version in ('reported-annual-3','reported-annual-4','reported-annual-5'), lipid_fixed_breaks=payload.get('reported',{}).get('map_scale_version')=='lipids-fixed-2021-2023-v1', include_glucose=version in ('reported-annual-4','reported-annual-5'), unified_rates=version=='reported-annual-5', terminology_version=payload.get('reported',{}).get('reported_semantics_version'), include_liver=payload.get('reported',{}).get('rate_policy',{}).get('contract_version')=='reported-recipient-rate-v3', include_renal='renal_contract' in payload.get('reported',{}))
    require(payload==expected, 'Site rebuild mismatch')
    return payload, report


def approve(root, id, reviewer, confirm_hash, data_rights=False, map_rights=False):
    require(reviewer.strip() and confirm_hash==id and data_rights and map_rights, 'Explicit named approval and rights review required')
    root = Path(root)
    with locked(root/'site'):
        payload, report = inspect(root,id)
        target = root/'site/releases'/id
        approval = dict(status='approved', release_id=id, data_sha256=id, reviewer=reviewer, approved_at=now(), validation=report,
                        data_rights_reviewed=True, map_rights_reviewed=True)
        if target.exists():
            require(digest((target/'data.json').read_bytes())==id, 'Corrupt immutable release')
            approval = read_json(target/'approval.json')
        else:
            atomic_write(target/'data.json',encoded(payload)); write_json(target/'approval.json',approval)
        write_json(root/'site/current.json',approval)
    return approval


def approved(root):
    root = Path(root); pointer = read_json(root/'site/current.json')
    id = pointer['release_id']
    require(len(id)==64 and all(c in '0123456789abcdef' for c in id), 'Invalid approved id')
    directory = checked_path(root,'site/releases/'+id)
    raw = (directory/'data.json').read_bytes()
    require(pointer.get('status')=='approved' and pointer.get('reviewer') and pointer==read_json(directory/'approval.json')
            and digest(raw)==id==pointer['data_sha256'], 'Invalid site approval')
    return pointer, raw


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('command',choices=['build','validate','approve'])
    p.add_argument('--data-dir',type=Path,default=Path('data'))
    p.add_argument('--analysis-release'); p.add_argument('--release-id')
    p.add_argument('--include-renal', action='store_true', help='Build a renal candidate only; does not approve or publish')
    p.add_argument('--reviewer',default=''); p.add_argument('--confirm-hash',default='')
    p.add_argument('--data-rights-reviewed',action='store_true'); p.add_argument('--map-rights-reviewed',action='store_true')
    a=p.parse_args()
    if a.command=='build': result=build(a.data_dir,a.analysis_release,include_renal=a.include_renal)
    elif a.command=='validate': result=inspect(a.data_dir,a.release_id)[1]
    else: result=approve(a.data_dir,a.release_id,a.reviewer,a.confirm_hash,a.data_rights_reviewed,a.map_rights_reviewed)
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__': main()
