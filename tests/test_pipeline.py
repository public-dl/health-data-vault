import copy
import io
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from hdv.collector import collect, discover, official_url, store_fetch
from hdv.common import CONFIG, digest, read_json, read_jsonl, write_jsonl
from hdv.normalizer import (MUNICIPAL_SCOPE, cell_state, extract, geography,
                            select_rows, structure)
from hdv.validator import comparable_series, validate_records, validate


def synthetic_records():
    """Independent synthetic populations: one municipality has all cases; 29 have zero."""
    indicators = read_json(CONFIG / "indicators.json")
    reference = read_json(CONFIG / "province_reference.json")
    records = []
    for year in (2021, 2022, 2023):
        for region in range(31):
            for indicator in indicators:
                column = indicator["column"]
                value = reference["years"][str(year)]["cells"][column + "6"] if region < 2 else 0
                records.append(dict(publication_fiscal_year=year+1, observation_fiscal_year=year,
                                    geography_name="県計" if region == 0 else f"試験{region}市",
                                    geography_level="prefecture_total" if region == 0 else "municipality",
                                    indicator_id=indicator["indicator_id"], original_label=indicator["indicator_id"],
                                    value=value, value_state="numeric" if value else "zero",
                                    unit=indicator["unit"], population_scope=MUNICIPAL_SCOPE,
                                    source_url="https://www.kenko-niigata.com/test.xlsx", source_sheet="test",
                                    source_cell=column+str(6 if region == 0 else region+6), source_row=6 if region == 0 else region+6,
                                    source_sha256="0"*64, comparability_status=indicator["comparability_status"],
                                    definition_version=f"niigata-tokutei-{year}-v0.1", validation_status="unvalidated",
                                    comparison_allowed=False, comparison_group=indicator["indicator_id"],
                                    included=True, selection_status="included", original_value=value, cached_value=value,
                                    source_data_type="n", number_format="General"))
    return records


class CellStateTests(unittest.TestCase):
    def test_distinct_states(self):
        for value, state in [(None, "blank"), ("", "empty_string"), ("-", "hyphen"), ("－", "hyphen"),
                             (0, "zero"), (1, "numeric"), ("#REF!", "excel_error"),
                             ("判定不能", "indeterminate_text"), ("未実施", "not_performed"),
                             ("未把握", "not_ascertained"), ("・", "not_applicable_marker"),
                             ("unknown", "unexpected_text"), (True, "unexpected_boolean")]:
            with self.subTest(value=value):
                w = Workbook(); c = w.active["A1"]; c.value = value
                result = cell_state(c, c)
                self.assertEqual(result["value_state"], state)
                self.assertEqual(result["original_value"], value)

    def test_display_hyphen_is_still_zero(self):
        w = Workbook(); c = w.active["A1"]; c.value = 0; c.number_format = '#,##0;-#,##0;"-"'
        self.assertEqual(cell_state(c, c)["value_state"], "zero")

    def test_formula_cache_absence_and_error(self):
        w = Workbook(); cached = Workbook(); c = w.active["A1"]; c.value = "=1/0"
        self.assertEqual(cell_state(c, cached.active["A1"])["value_state"], "formula_cache_missing")
        cached.active["A1"] = "#DIV/0!"
        result = cell_state(c, cached.active["A1"])
        self.assertEqual(result["value_state"], "excel_error")
        self.assertEqual(result["formula"], "=1/0")


class CollectorTests(unittest.TestCase):
    def test_relative_link_and_duplicate_anchor(self):
        html = '<a href="/material/16_R3_tokuteikenshin.xlsx">対象</a>' * 2
        self.assertEqual(discover(html, "https://www.kenko-niigata.com/a/page.html", "16_R3_tokuteikenshin.xlsx"),
                         "https://www.kenko-niigata.com/material/16_R3_tokuteikenshin.xlsx")

    def test_missing_ambiguous_or_external_link(self):
        for html in ["<p>no file</p>", '<a href="/a/test.xlsx"></a><a href="/b/test.xlsx"></a>',
                     '<a href="https://example.com/test.xlsx"></a>']:
            with self.subTest(html=html), self.assertRaises(ValueError):
                discover(html, "https://www.kenko-niigata.com/page.html", "test.xlsx")

    def test_unapproved_urls(self):
        for url in ["http://www.kenko-niigata.com/a", "https://example.com/a", "file:///a",
                    "https://user@www.kenko-niigata.com/a", "https://www.kenko-niigata.com.evil.test/a"]:
            with self.subTest(url=url), self.assertRaises(ValueError):
                official_url(url)

    def test_content_addressing_and_change_history(self):
        with tempfile.TemporaryDirectory() as tmp:
            url = "https://www.kenko-niigata.com/test.xlsx"
            a = store_fetch(tmp, url, b"original", {}, url, "excel")
            b = store_fetch(tmp, url, b"original", {}, url, "excel")
            c = store_fetch(tmp, url, b"updated", {}, url, "excel")
            self.assertFalse(a["content_changed"]); self.assertFalse(b["content_changed"])
            self.assertTrue(c["content_changed"]); self.assertEqual(c["previous_sha256"], a["sha256"])
            self.assertEqual((Path(tmp)/a["raw_path"]).read_bytes(), b"original")
            self.assertEqual(len(list((Path(tmp)/"raw/events").glob("*.json"))), 3)

    def test_corrupt_existing_object_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as tmp:
            url = "https://www.kenko-niigata.com/test.xlsx"
            a = store_fetch(tmp, url, b"original", {}, url, "excel")
            (Path(tmp)/a["raw_path"]).write_bytes(b"corrupted")
            with self.assertRaises(ValueError):
                store_fetch(tmp, url, b"original", {}, url, "excel")

    def test_collector_fetches_page_before_discovered_excel(self):
        contracts = read_json(CONFIG/"contracts.json")
        calls = []
        def fetcher(url):
            calls.append(url)
            if url.endswith(".html"):
                contract = next(c for c in contracts.values() if c["page_url"] == url)
                data = f'<a href="/changed/path/{contract["filename"]}">Excel</a>'.encode()
            else:
                w = Workbook(); stream = io.BytesIO(); w.save(stream); data = stream.getvalue()
            return data, {}, url
        with tempfile.TemporaryDirectory() as tmp:
            manifest = collect(tmp, fetcher=fetcher)
            self.assertEqual(len(calls), 6)
            self.assertEqual(len(manifest["sources"]), 3)
            self.assertTrue(all("/changed/path/" in s["source_url"] for s in manifest["sources"]))


class NormalizerTests(unittest.TestCase):
    def test_geography_levels(self):
        expected = {"県計": "prefecture_total", "市　 計": "city_total", "町村計": "town_village_total",
                    "村上保健所計": "health_center_total", "新潟市": "municipality",
                    "関川村": "municipality", "国保組合計": "insurance_union_total", "医師国保": "insurance_union"}
        for name, level in expected.items():
            self.assertEqual(geography(name)[1], level)

    def pair(self):
        records = synthetic_records()[:41]
        for r in records:
            r.update(geography_name="新潟市", geography_level="municipality", source_row=50)
        duplicate = copy.deepcopy(records)
        for r in duplicate:
            r["source_row"] = 51
        return records+duplicate

    def test_documented_duplicate_preserved_with_reason(self):
        records = self.pair(); decisions, issues = select_rows(records)
        self.assertEqual(len(records), 82)
        self.assertEqual(sum(r["included"] for r in records), 41)
        self.assertEqual(decisions[0]["retained_row"], 50)
        self.assertEqual(issues[0]["severity"], "warning")

    def test_conflicting_duplicate_quarantines_both(self):
        records = self.pair(); records[-1]["original_value"] += 1
        decisions, issues = select_rows(records)
        self.assertFalse(any(r["included"] for r in records))
        self.assertIsNone(decisions[0]["retained_row"])
        self.assertEqual(issues[0]["severity"], "error")

    def test_unknown_identical_duplicate_is_not_silently_resolved(self):
        records = self.pair()
        for r in records:
            r["geography_name"] = "村上市"
        select_rows(records)
        self.assertFalse(any(r["included"] for r in records))

    def test_structure_changes_block_all_records_but_archive_unknown_cells(self):
        for mutation in ("column", "header", "merge", "sheet"):
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory() as tmp:
                w = Workbook(); w.active.title = "overview"; s = w.create_sheet("detail ")
                s["A6"] = "県計"; s["B4"] = "受診者数"; s["B6"] = 1
                contract = {"structure": structure(w), "main_sheet": s.title}
                if mutation == "column": s["AQ4"] = "未知列"
                if mutation == "header": s["B4"] = "変更された定義"
                if mutation == "merge": s.merge_cells("C3:D3")
                if mutation == "sheet": s.title = "detail"
                path = Path(tmp)/"fixture.xlsx"; w.save(path)
                source = dict(observation_fiscal_year=2021, publication_fiscal_year=2022,
                              source_url="https://www.kenko-niigata.com/a.xlsx", sha256=digest(path.read_bytes()),
                              annual_page_url="https://www.kenko-niigata.com/a.html", retrieved_at="2026-09-17T00:00:00+09:00")
                records, cells, issues, _ = extract(path, source, contract, read_json(CONFIG/"indicators.json"))
                self.assertIn("structure_changed", [i["code"] for i in issues])
                self.assertFalse(any(r["included"] for r in records))
                self.assertTrue(cells)
                if mutation == "column": self.assertTrue(any(c["source_cell"] == "AQ4" for c in cells))


class ValidatorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.records = synthetic_records()

    def test_123_controls_and_partition_totals(self):
        records, report = validate_records(self.records)
        self.assertEqual(report["errors"], 0)
        self.assertEqual(report["province_regression"]["passed"], 123)
        self.assertTrue(any(r["comparison_allowed"] for r in records))

    def test_regression_mismatch_fails_closed(self):
        records = copy.deepcopy(self.records); records[0]["value"] += 1
        result, report = validate_records(records)
        self.assertIn("province_regression", [i["code"] for i in report["issues"]])
        self.assertFalse(any(r["comparison_allowed"] for r in result))

    def test_duplicate_year_indicator_geography_key(self):
        _, report = validate_records(self.records + [self.records[0]])
        self.assertIn("duplicate_key", [i["code"] for i in report["issues"]])

    def test_unknown_indicator_year_unit_error(self):
        for field, value, code in [("indicator_id", "new", "unknown_indicator"),
                                   ("publication_fiscal_year", 2025, "unexpected_year"),
                                   ("unit", "%", "unit_changed"), ("value_state", "excel_error", "excel_error")]:
            with self.subTest(field=field):
                records = copy.deepcopy(self.records); records[0][field] = value
                _, report = validate_records(records)
                self.assertIn(code, [i["code"] for i in report["issues"]])

    def test_missing_is_not_zero_or_eligible(self):
        records = copy.deepcopy(self.records)
        target = next(r for r in records if r["indicator_id"] == "metabo_indeterminate" and r["geography_name"] == "試験2市")
        target.update(value=None, value_state="blank")
        result, report = validate_records(records)
        self.assertEqual(report["errors"], 0)
        preserved = next(r for r in result if r["source_cell"] == target["source_cell"] and r["observation_fiscal_year"] == 2021)
        self.assertIsNone(preserved["value"])
        self.assertFalse(preserved["comparison_allowed"])

    def test_indeterminate_count_is_a_numeric_category(self):
        result, _ = validate_records(self.records)
        r = next(r for r in result if r["indicator_id"] == "metabo_indeterminate")
        self.assertEqual(r["value"], 509)
        self.assertTrue(r["comparison_allowed"])

    def test_comparison_guard(self):
        result, _ = validate_records(self.records)
        for id in ("metabo_case", "hba1c", "ecg"):
            series = [r for r in result if r["indicator_id"] == id and r["geography_name"] == "県計"]
            if id == "metabo_case":
                self.assertEqual([r["observation_fiscal_year"] for r in comparable_series(series)], [2021, 2022, 2023])
                with self.assertRaises(ValueError): comparable_series(series+[series[0]])
            else:
                with self.assertRaises(ValueError): comparable_series(series)
                # Changing only the data flags cannot promote a prohibited indicator.
                for r in series: r.update(comparability_status="compatible", comparison_allowed=True)
                with self.assertRaises(ValueError): comparable_series(series)

    def test_unvalidated_cannot_be_compared(self):
        r = next(r for r in self.records if r["indicator_id"] == "metabo_case")
        with self.assertRaises(ValueError): comparable_series([r])


@unittest.skipUnless(Path("data/processed/latest.json").exists(), "Run python -m hdv run first for official-data regression")
class OfficialDataTests(unittest.TestCase):
    def copy_run(self, root):
        original = Path("data/processed") / read_json("data/processed/latest.json")["run_id"]
        output = root / "processed/test"
        output.mkdir(parents=True)
        for name in ("collection.json", "run.json", "records.jsonl"):
            shutil.copyfile(original/name, output/name)
        for source in read_json(output/"collection.json")["sources"]:
            for relative in (source["raw_path"], "raw/objects/"+source["page_sha256"]+".html"):
                dest = root/relative; dest.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(Path("data")/relative, dest)
        return output

    def test_modified_normalized_values_are_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); output = self.copy_run(root)
            records = read_jsonl(output/"records.jsonl"); records[0]["value"] += 1
            write_jsonl(output/"records.jsonl", records)
            report = validate(root, output)
            self.assertIn("normalized_records_differ_from_raw", [i["code"] for i in report["issues"]])
            self.assertEqual(read_jsonl(output/"comparable_records.jsonl"), [])

    def test_failed_revalidation_clears_old_comparison_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); output = self.copy_run(root)
            write_jsonl(output/"comparable_records.jsonl", [{"old": "approved"}])
            source = read_json(output/"collection.json")["sources"][0]
            (root/source["raw_path"]).unlink()
            with self.assertRaises(FileNotFoundError): validate(root, output)
            self.assertEqual(read_jsonl(output/"comparable_records.jsonl"), [])
            self.assertEqual(read_json(output/"validation_report.json")["status"], "incomplete")

    def test_official_result_and_preserved_raw(self):
        output = Path("data/processed") / read_json("data/processed/latest.json")["run_id"]
        report = read_json(output/"validation_report.json")
        self.assertEqual(report["errors"], 0)
        self.assertEqual(report["province_regression"]["passed"], 123)
        records = read_jsonl(output/"validated_records.jsonl")
        self.assertEqual(len(records), 5986)
        self.assertEqual(sum(r["selection_status"] == "duplicate_excluded" for r in records), 123)
        self.assertEqual(sum(r["selection_status"] == "out_of_scope" for r in records), 328)
        for source in read_json(output/"collection.json")["sources"]:
            self.assertEqual(digest((Path("data")/source["raw_path"]).read_bytes()), source["sha256"])
        comparable = read_jsonl(output/"comparable_records.jsonl")
        self.assertEqual(len(comparable), 1074)
        self.assertTrue(all(r["comparability_status"] == "compatible" and r["included"] for r in comparable))


if __name__ == "__main__":
    unittest.main()
